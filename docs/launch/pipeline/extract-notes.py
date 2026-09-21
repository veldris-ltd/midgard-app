"""Write the deck's speaker notes to a Markdown run-sheet, one section per slide."""
import sys
from pptx import Presentation

deck, out = sys.argv[1], sys.argv[2]
p = Presentation(deck)
lines = ["# Midgard launch — speaker run-sheet", "",
         "Generated from the speaker notes in `Midgard-Launch.pptx`. Timings in brackets add up to roughly 18 minutes; the demo video is 4 min 09 s.", ""]
for i, slide in enumerate(p.slides, 1):
    title = None
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip().split("\n")[0]
            if len(t) > 12 and not t.startswith(("0", "1", "MIDGARD ·")):
                title = t
                break
    notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
    lines.append(f"## {i:02d} · {title or 'Title'}")
    lines.append("")
    lines.append(notes or "_No notes._")
    lines.append("")
open(out, "w", encoding="utf-8").write("\n".join(lines))
print("wrote", out, len(p.slides), "slides")
