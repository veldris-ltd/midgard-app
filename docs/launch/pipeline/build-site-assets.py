"""Prepare the launch site's media: slide images, thumbnails, posters, decks, videos, fonts, notes.

Run from the pipeline directory after both decks and videos exist:
    python build-site-assets.py <site-dir>
Reads ./out (dark) and ./out-light (light) when a pipeline build is present, otherwise the
committed decks, videos and slide PNGs in the kit folder above. Writes <site-dir>/media and
<site-dir>/assets. The .mp4 and .pptx copies under <site-dir>/media are not committed.
"""
import json, os, shutil, subprocess, sys
from PIL import Image
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(HERE, 'site')
KIT = os.path.join(HERE, '..')
REPO_FONTS = os.environ.get('MG_FONTS') or os.path.join(HERE, '..', '..', '..', 'ui', 'design-system', 'src', 'fonts')

MODES = {
    'dark':  dict(out='out',       deck='Midgard-Launch.pptx',       deck_as='Midgard-Launch.pptx',       video='midgard-launch-demo.mp4',       video_as='midgard-launch-demo.mp4'),
    'light': dict(out='out-light', deck='Midgard-Launch-light.pptx', deck_as='Midgard-Launch_light.pptx', video='midgard-launch-demo-light.mp4', video_as='midgard-launch-demo-light.mp4'),
}
TITLES = ['Midgard', 'Agenda', 'The problem', 'The suite', 'Design principles', 'Architecture', 'Object fabric',
          'Governance', 'Forge', 'Mimir', 'Heimdall', 'Bifrost', 'Admin and clients', 'Live demo',
          'Post-quantum by default', 'Proof', 'Where we are', 'Next steps']


def ffmpeg():
    exe = shutil.which('ffmpeg')
    if exe:
        return exe
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()


def duration(path):
    p = subprocess.run([ffmpeg(), '-hide_banner', '-i', path], capture_output=True, text=True)
    line = next(l for l in p.stderr.splitlines() if 'Duration' in l)
    h, m, s = line.split('Duration:')[1].split(',')[0].strip().split(':')
    return int(h) * 3600 + int(m) * 60 + float(s)


manifest = {'slides': len(TITLES), 'titles': TITLES, 'modes': {}}
for mode, c in MODES.items():
    # Prefer a fresh pipeline build (./out, ./out-light); otherwise use the committed kit one level up.
    built = os.path.isdir(os.path.join(HERE, c['out'], 'png'))
    src = os.path.join(HERE, c['out']) if built else KIT
    png = os.path.join(src, 'png') if built else os.path.join(KIT, 'slides' if mode == 'dark' else 'slides-light')
    deck_src = os.path.join(src, c['deck'] if built else c['deck_as'])
    dst = os.path.join(SITE, 'media', mode)
    os.makedirs(os.path.join(dst, 'slides'), exist_ok=True)
    os.makedirs(os.path.join(dst, 'thumbs'), exist_ok=True)
    for i in range(1, len(TITLES) + 1):
        im = Image.open(os.path.join(png, f'Slide{i}.PNG')).convert('RGB')
        im.resize((1600, 900), Image.LANCZOS).save(os.path.join(dst, 'slides', f'{i:02d}.webp'), quality=86, method=6)
        im.resize((320, 180), Image.LANCZOS).save(os.path.join(dst, 'thumbs', f'{i:02d}.webp'), quality=78, method=6)
    shutil.copyfile(deck_src, os.path.join(dst, c['deck_as']))
    shutil.copyfile(os.path.join(src, c['video']), os.path.join(dst, c['video_as']))
    poster = os.path.join(dst, 'poster.jpg')
    subprocess.run([ffmpeg(), '-y', '-hide_banner', '-loglevel', 'error', '-ss', '4', '-i', os.path.join(src, c['video']),
                    '-frames:v', '1', '-q:v', '3', poster], check=True)
    secs = duration(os.path.join(src, c['video']))
    manifest['modes'][mode] = {
        'deck': c['deck_as'], 'deckBytes': os.path.getsize(os.path.join(dst, c['deck_as'])),
        'video': c['video_as'], 'videoBytes': os.path.getsize(os.path.join(dst, c['video_as'])),
        'videoSeconds': round(secs),
    }

# Speaker notes are identical in both decks; read them once.
deck = Presentation(os.path.join(SITE, 'media', 'dark', MODES['dark']['deck_as']))
manifest['notes'] = [s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else '' for s in deck.slides]
# Presenter-only fallback advice makes no sense on the web page.
import re
manifest['notes'] = [re.sub(r"\s*If the embedded video will not play, open \S+ beside the deck\.", "", n) for n in manifest['notes']]

os.makedirs(os.path.join(SITE, 'assets', 'fonts'), exist_ok=True)
for f in os.listdir(REPO_FONTS):
    if f.endswith(('.ttf', '-OFL.txt')):
        shutil.copyfile(os.path.join(REPO_FONTS, f), os.path.join(SITE, 'assets', 'fonts', f))
# Social card: the dark title slide.
Image.open(os.path.join(HERE, 'out', 'png', 'Slide1.PNG') if os.path.isdir(os.path.join(HERE, 'out', 'png')) else os.path.join(KIT, 'slides', 'Slide1.PNG')).convert('RGB').resize((1200, 675), Image.LANCZOS).save(
    os.path.join(SITE, 'assets', 'og.jpg'), quality=88)
# A script rather than JSON, so the page also works when opened straight from disk.
open(os.path.join(SITE, 'assets', 'kit.js'), 'w', encoding='utf-8').write('window.KIT = ' + json.dumps(manifest, ensure_ascii=False, indent=1) + ';' + chr(10))
print(json.dumps(manifest['modes'], indent=1))
